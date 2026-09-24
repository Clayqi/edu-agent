"""PPT 优化（「边框里的 PPT 优化」一期）：体检 → 保守原地统一 → 内容零丢失校验。

设计要点（对应 `docs/17-边框里的PPT优化计划书.md` §4.1.1）：

* **确定性优先**：体检 / 台账 / 校验三件交给 ppt-master 的确定性脚本
  （`pptx_intake.py`、`beautify_inventory.py`）；"重排"那种模型驱动的活不在本模块里。
* **原地统一自己写**：`统一字体`（含 `<a:ea>` 中文字体）、`统一字号层级`、`清理动画/转场`
  用 python-pptx + lxml 直接改 XML —— ppt-master 的 beautify 档是"重排"，不做这种保守补丁。
* **原稿只读**：所有产物写到 `data/tmp/ppt_jobs/<job_id>/`，原文件从不修改；
  用户点下载后再复制到 `data/ppt_out/`。
* **内容零丢失**：改完可跑 `verify()`（= ppt-master 的 `beautify_inventory.py --verify`），
  缺一个原串就判失败；同时报"多出来的字符数"。
* **未安装可降级**：ppt-master 不在时 `status()` 报清楚缺什么，`intake()`/`verify()` 明确失败，
  但 `polish()` 的原地动作仍可用（只依赖 python-pptx）。
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import time
import uuid
from pathlib import Path

# ---------- 安装位置 ----------

_ENV_HOME = "PPT_MASTER_HOME"
_DEFAULT_HOME = Path(os.environ.get("LOCALAPPDATA", "")) / "edu-agent" / "deps" / "ppt-master"

# 一期用到的确定性脚本（缺任何一个，"体检/校验"就不可用）
_REQUIRED_SCRIPTS = ("pptx_intake.py", "beautify_inventory.py", "beautify_identity.py")
# 二期（模型驱动重排）会用到的，只做存在性上报
_OPTIONAL_SCRIPTS = ("pptx_to_svg.py", "svg_to_pptx.py", "pptx_template_import.py",
                     "svg_quality_checker.py", "pptx_transitions.py", "pptx_animations.py",
                     "pptx_delivery_check.py", "project_manager.py")
# 一期需要的第三方库（缺了会影响对应能力，但多数路径仍可用）
_PY_DEPS = (("pptx", "python-pptx", "生成/改写 pptx"),
            ("lxml", "lxml", "改 XML（python-pptx 依赖）"),
            ("yaml", "PyYAML", "模板注册"),
            ("xlsxwriter", "XlsxWriter", "导出时的数据表"),
            ("PIL", "Pillow", "图片尺寸优化"),
            ("numpy", "numpy", "图片处理"))


def home() -> Path:
    """ppt-master 根目录：环境变量 > %LOCALAPPDATA%\\edu-agent\\deps\\ppt-master。"""
    env = os.environ.get(_ENV_HOME)
    return Path(env) if env else _DEFAULT_HOME


def scripts_dir() -> Path:
    return home() / "skills" / "ppt-master" / "scripts"


def workroot() -> Path:
    """任务临时目录（仓库内，data/ 已被 gitignore）。"""
    root = Path(__file__).resolve().parents[2] / "data" / "tmp" / "ppt_jobs"
    root.mkdir(parents=True, exist_ok=True)
    return root


def outdir() -> Path:
    d = Path(__file__).resolve().parents[2] / "data" / "ppt_out"
    d.mkdir(parents=True, exist_ok=True)
    return d


# ---------- 自检 ----------

def _commit(root: Path) -> str:
    try:
        r = subprocess.run(["git", "-C", str(root), "rev-parse", "--short", "HEAD"],
                           capture_output=True, text=True, timeout=20)
        return r.stdout.strip() if r.returncode == 0 else ""
    except Exception:  # noqa: BLE001 自检不该抛
        return ""


def status() -> dict:
    """安装自检：脚本齐不齐、依赖缺不缺、能不能干活。"""
    root, sd = home(), scripts_dir()
    present = {name: (sd / name).exists() for name in _REQUIRED_SCRIPTS + _OPTIONAL_SCRIPTS}
    missing_req = [n for n in _REQUIRED_SCRIPTS if not present[n]]
    deps = {}
    for mod, pkg, why in _PY_DEPS:
        try:
            __import__(mod)
            deps[pkg] = {"ok": True, "use": why}
        except Exception:  # noqa: BLE001
            deps[pkg] = {"ok": False, "use": why}
    missing_dep = [p for p, v in deps.items() if not v["ok"]]
    return {
        "ok": not missing_req and not missing_dep,
        "installed": root.is_dir() and (sd / "pptx_intake.py").exists(),
        "root": str(root),
        "commit": _commit(root) if root.is_dir() else "",
        "scripts": present,
        "missing_required_scripts": missing_req,
        "python_deps": deps,
        "missing_python_deps": missing_dep,
        # 一期只用这三个能力；缺可选脚本只影响二期
        "intake_ready": not missing_req,
        "note": ("一期用 pptx_intake + beautify_inventory（体检/台账/校验）；"
                 "二期的 SVG 重排还需要强模型，见 docs/17 §4.1.1"),
    }


# ---------- 任务目录 ----------

def new_job(src: Path) -> dict:
    """把上传的 .pptx 拷进任务目录（原文件只读），返回 job 信息。"""
    jid = time.strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:6]
    jd = workroot() / jid
    (jd / "analysis").mkdir(parents=True, exist_ok=True)
    dst = jd / "source.pptx"
    shutil.copy2(src, dst)
    return {"job_id": jid, "dir": str(jd), "source": str(dst), "name": Path(src).name}


def job_dir(job_id: str) -> Path:
    return workroot() / job_id


def cleanup_job(job_id: str) -> None:
    shutil.rmtree(job_dir(job_id), ignore_errors=True)


def cleanup_stale(max_age_hours: float = 24.0) -> int:
    """清掉超过 max_age_hours 的任务目录（启动时调）。"""
    n, now = 0, time.time()
    for d in workroot().glob("*"):
        try:
            if d.is_dir() and (now - d.stat().st_mtime) > max_age_hours * 3600:
                shutil.rmtree(d, ignore_errors=True)
                n += 1
        except OSError:
            pass
    return n


# ---------- 调 ppt-master 的确定性脚本 ----------

def _run(args: list[str], timeout: int = 600) -> dict:
    """跑一个 ppt-master 脚本，返回 {ok, code, out, err}（不抛）。"""
    py = sys.executable
    try:
        r = subprocess.run([py] + args, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=timeout)
        return {"ok": r.returncode == 0, "code": r.returncode,
                "out": (r.stdout or "").strip(), "err": (r.stderr or "").strip()}
    except subprocess.TimeoutExpired:
        return {"ok": False, "code": -1, "out": "", "err": f"超时（>{timeout}s）"}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "code": -1, "out": "", "err": f"{type(e).__name__}: {e}"}


def _read_json(p: Path) -> dict:
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def _pick_font(block: dict) -> str:
    """从主题字体块里取"实际会用的"字体名：优先 ea（中文），其次 latin，再取 Hans 补充面。"""
    if not isinstance(block, dict):
        return ""
    return (block.get("ea") or block.get("latin")
            or (block.get("scripts") or {}).get("Hans") or "")


def intake(job: dict) -> dict:
    """体检：拆包 + 逐页台账，返回精简报告（给面板用）。

    产物落在 `<job>/analysis/`：`source_profile.json`、`*.slide_library.json`、
    `*.identity.json`、`beautify_inventory.json` —— `verify()` 复用它们。
    """
    st = status()
    if not st["intake_ready"]:
        return {"ok": False, "error": "ppt-master 未就绪", "status": st,
                "hint": "跑 deploy/setup_ppt_master.cmd，或设 PPT_MASTER_HOME 指到已安装目录"}

    sd, jd = scripts_dir(), Path(job["dir"])
    adir = jd / "analysis"
    src = Path(job["source"])

    r1 = _run([str(sd / "pptx_intake.py"), "-o", str(adir), str(src)])
    if not r1["ok"]:
        return {"ok": False, "error": "pptx_intake 失败", "detail": r1["err"] or r1["out"]}

    prof = _read_json(adir / "source_profile.json")
    decks = prof.get("decks") or []
    deck = decks[0] if decks else {}
    libs = sorted(adir.glob("*.slide_library.json"))
    if not libs:
        return {"ok": False, "error": "没找到 slide_library.json（不是标准 pptx？）"}

    inv = adir / "beautify_inventory.json"
    r2 = _run([str(sd / "beautify_inventory.py"), str(libs[0]), "-o", str(inv)])
    if not r2["ok"]:
        return {"ok": False, "error": "beautify_inventory 失败", "detail": r2["err"] or r2["out"]}
    r3 = _run([str(sd / "beautify_inventory.py"), str(inv), "--summary"])
    summ = {}
    if r3["ok"]:
        try:
            summ = json.loads(r3["out"])
        except Exception:  # noqa: BLE001
            summ = {}

    ident = _read_json(next(iter(sorted(adir.glob("*.identity.json"))), Path("nonexistent")))
    theme = (ident.get("theme") or {}) if ident else {}
    tfonts = theme.get("fonts") or {}
    sizes = (ident.get("sizes") or theme.get("sizes") or {}) if ident else {}
    observed = ident.get("observed") or {}

    rows = []
    for s in summ.get("slides") or []:
        rows.append({
            "page": s.get("slide_index"),
            "type": s.get("page_type", ""),
            "blocks": s.get("text_block_count", 0),
            "chars": s.get("text_char_count", 0),
            "tables": s.get("table_count", 0),
            "charts": s.get("chart_count", 0),
            "diagrams": s.get("diagram_count", 0),
            "images": s.get("image_count", 0),
            "ignored": s.get("ignored") or [],
            "confirm": s.get("needs_confirmation") or [],
        })
    chars = [r["chars"] for r in rows if r["chars"]]
    dense = [r["page"] for r in rows if r["chars"] and r["chars"] > (max(chars) if chars else 0) * 0.9] if chars else []
    sparse = [r["page"] for r in rows if 0 < r["chars"] < 20]

    return {
        "ok": True,
        "job_id": job["job_id"],
        "file": job["name"],
        "slides": deck.get("slide_count") or len(rows),
        "canvas": deck.get("canvas") or ident.get("canvas") or {},
        "theme": {
            "title_font": _pick_font(tfonts.get("title")),
            "body_font": _pick_font(tfonts.get("body")),
            "latin_title": (tfonts.get("title") or {}).get("latin", ""),
            "latin_body": (tfonts.get("body") or {}).get("latin", ""),
            "title_size": sizes.get("title"),
            "body_size": sizes.get("body"),
            "body_levels": sizes.get("body_levels") or [],
            "palette": ident.get("theme_palette") or (theme.get("palette") or {}),
        },
        "observed": {
            "fonts_latin": observed.get("fonts", {}).get("latin") or [],
            "fonts_ea": observed.get("fonts", {}).get("ea") or [],
            "sizes_pt": observed.get("sizes_pt") or [],
            "colors": observed.get("colors") or [],
        },
        "totals": {
            "tables": sum(r["tables"] for r in rows),
            "charts": sum(r["charts"] for r in rows),
            "diagrams": sum(r["diagrams"] for r in rows),
            "images": sum(r["images"] for r in rows),
            "chars": sum(r["chars"] for r in rows),
        },
        "pages": rows,
        "flags": {
            "dense_pages": dense,
            "sparse_pages": sparse,
            "needs_confirmation": sorted({p["page"] for r in rows for p in [r] if r["confirm"]}),
        },
        "artifacts": {"dir": str(jd), "analysis": str(adir), "inventory": str(inv),
                      "slide_library": str(libs[0])},
    }


def verify(job: dict, exported: Path) -> dict:
    """内容零丢失校验：每个原串都得在它那一页上；并报"多出来的字符数"。"""
    inv = Path(job["dir"]) / "analysis" / "beautify_inventory.json"
    if not inv.exists():
        return {"ok": False, "error": "还没有台账，先体检"}
    r = _run([str(scripts_dir() / "beautify_inventory.py"), str(inv), "--verify", str(exported)])
    return {"ok": r["ok"], "detail": r["out"] or r["err"],
            "note": "ok=true 表示逐页逐串都在（内容没被改掉）"}


# ---------- 保守原地统一（我们自己实现；ppt-master 不做这个） ----------

_ACTION_LABELS = {
    "font": "统一字体（含中文字体 <a:ea>）",
    "size": "统一字号层级",
    "strip-anim": "清理动画 / 转场",
}


def _iter_text_frames(prs):
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield idx, shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield idx, cell.text_frame


def _set_ea_font(run, name: str) -> None:
    """给 run 补 `<a:ea typeface="…">`（python-pptx 不直接管东亚字体）。"""
    rPr = run._r.get_or_add_rPr()  # noqa: SLF001 python-pptx 没暴露，只能这样
    for tag in ("a:ea",):
        el = rPr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag.split(':')[1]}")
        if el is None:
            from lxml import etree  # noqa: PLC0415
            el = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        el.set("typeface", name)


def _action_font(prs, target_ea: str, target_latin: str) -> dict:
    runs = 0
    for _page, tf in _iter_text_frames(prs):
        for para in tf.paragraphs:
            for run in para.runs:
                if not (run.text or "").strip():
                    continue
                run.font.name = target_latin
                _set_ea_font(run, target_ea)
                runs += 1
    return {"runs": runs}


def _action_size(prs, anchors: list[float], tolerance: float = 1.0) -> dict:
    """把散落的字号归并到最近的层级锚（保留相对层级，不做等比缩放）。"""
    from pptx.util import Pt  # noqa: PLC0415
    changed, seen = 0, {}
    for _page, tf in _iter_text_frames(prs):
        for para in tf.paragraphs:
            for run in para.runs:
                size = run.font.size
                if size is None:
                    continue
                pt = float(size.pt)
                near = min(anchors, key=lambda a: abs(a - pt))
                seen[pt] = seen.get(pt, 0) + 1
                if abs(near - pt) > tolerance:
                    run.font.size = Pt(near)
                    changed += 1
    return {"changed": changed, "observed_sizes": sorted(seen.items(), key=lambda x: -x[1])}


def _action_strip_anim(prs, strip_transition: bool = True) -> dict:
    """删每页 `p:timing`（动画）与可选 `p:transition`（切页效果）。"""
    from lxml import etree  # noqa: PLC0415
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    timing = transition = 0
    for slide in prs.slides:
        el = slide._element  # noqa: SLF001
        for node in el.findall(f"{P}timing"):
            el.remove(node)
            timing += 1
        if strip_transition:
            for node in list(el):
                if node.tag in (f"{P}transition",
                                "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
                    if node.tag == f"{P}transition":
                        el.remove(node)
                        transition += 1
    return {"timing_removed": timing, "transition_removed": transition}


def polish(job: dict, actions: list[str], *, target_ea: str = "", target_latin: str = "",
           size_anchors: list[float] | None = None) -> dict:
    """在任务目录里对副本做**保守原地统一**，返回新文件路径与逐项报告。

    原稿不动；产物 `<job>/polished.pptx`。动作：font / size / strip-anim。
    """
    from pptx import Presentation  # noqa: PLC0415

    jd = Path(job["dir"])
    src, dst = Path(job["source"]), jd / "polished.pptx"
    if not src.exists():
        return {"ok": False, "error": "任务目录里没有源文件"}
    actions = [a for a in (actions or []) if a in _ACTION_LABELS]
    if not actions:
        return {"ok": False, "error": "没选任何动作", "available": list(_ACTION_LABELS)}

    prs = Presentation(str(src))
    report: dict = {"ok": True, "applied": [], "pages": len(prs.slides.__iter__.__self__._sldIdLst)}  # noqa: SLF001
    if "font" in actions:
        ea = target_ea or "微软雅黑"
        latin = target_latin or "Calibri"
        got = _action_font(prs, ea, latin)
        got.update({"target_ea": ea, "target_latin": latin})
        report["applied"].append({"action": "font", "label": _ACTION_LABELS["font"], **got})
    if "size" in actions:
        anchors = sorted(size_anchors or [44.0, 32.0, 28.0, 24.0, 20.0, 18.0, 16.0, 14.0, 12.0])
        got = _action_size(prs, anchors)
        got["anchors"] = anchors
        report["applied"].append({"action": "size", "label": _ACTION_LABELS["size"], **got})
    if "strip-anim" in actions:
        got = _action_strip_anim(prs)
        report["applied"].append({"action": "strip-anim", "label": _ACTION_LABELS["strip-anim"], **got})

    prs.save(str(dst))
    report["file"] = str(dst)
    report["size_kb"] = round(dst.stat().st_size / 1024, 1)
    return report


def selftest(sample: Path | None = None) -> dict:
    """命令行自检：拿一份真 pptx 跑 体检 → 统一字体+清动画 → 内容校验。"""
    sample = sample or (Path(__file__).resolve().parents[2] / "content" / "plans" /
                        "完整课时模板（默认）.pptx")
    out: dict = {"status": status(), "sample": str(sample), "exists": sample.exists()}
    if not sample.exists():
        return out
    job = new_job(sample)
    out["job"] = job["job_id"]
    rep = intake(job)
    out["intake_ok"] = rep.get("ok")
    out["intake_brief"] = {k: rep.get(k) for k in ("slides", "totals", "flags", "theme")} if rep.get("ok") else rep
    pol = polish(job, ["font", "strip-anim"])
    out["polish"] = pol
    if pol.get("ok"):
        out["verify"] = verify(job, Path(pol["file"]))
    return out


if __name__ == "__main__":  # pragma: no cover - 手工自检入口
    if "--cleanup" in sys.argv:
        print("清理过期任务目录：", cleanup_stale(), "个")
    else:
        r = selftest()
        print(json.dumps(r, ensure_ascii=False, indent=2)[:4000])
