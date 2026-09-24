# -*- coding: utf-8 -*-
"""PPT 优化（一期）离线测试：安装自检 / 能力开关与门禁 / 任务目录 / 保守原地统一 / 体检报告解析。

设计：**默认全离线** —— 不联网、不调模型、不需要真的装 ppt-master：
- 安装自检与"未装就降级"直接用临时目录造出各种缺失场景；
- 体检报告用**假脚本**（临时目录里放两个 stub）验解析逻辑，不依赖真安装；
- `polish()` 用 python-pptx 现造一份小稿验动作，并断言**原稿字节不变**。

用法：
    python -m unittest discover -s tests -p "test_*.py"
"""
import json
import os
import sys
import tempfile
import time
import unittest
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import capabilities as cap          # noqa: E402
from edu_agent import ppt_polish as pp             # noqa: E402

os.environ.setdefault("EDU_DATA_DIR", tempfile.mkdtemp(prefix="edu_test_data_"))  # noqa: E402
from edu_agent import web_server as ws             # noqa: E402

_REQ = ("pptx_intake.py", "beautify_inventory.py", "beautify_identity.py")


def _make_pptx(path: Path, *, with_anim: bool = False) -> None:
    """造一份小稿：2 页，含显式字体/字号与（可选）动画节点。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "第一节 幂函数"
    r.font.name = "宋体"
    r.font.size = Pt(23)                     # 故意给个非锚点字号，验"归并"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    r2 = tb2.text_frame.paragraphs[0].add_run()
    r2.text = "正文内容"
    r2.font.name = "Arial"
    r2.font.size = Pt(19)
    if with_anim:
        from lxml import etree
        P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
        etree.SubElement(s2._element, f"{P}timing")      # noqa: SLF001 造个动画节点
    prs.save(str(path))


class _Sandbox(unittest.TestCase):
    """把 PPT_MASTER_HOME / EDU_DATA_DIR 指到临时目录，绝不碰真实安装与真实状态。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._old = {k: os.environ.get(k) for k in ("EDU_DATA_DIR", "PPT_MASTER_HOME")}
        os.environ["EDU_DATA_DIR"] = self._tmp.name
        os.environ["PPT_MASTER_HOME"] = str(Path(self._tmp.name) / "ppt-master")
        cap._CACHE.clear()
        self.addCleanup(self._restore)

    def _restore(self):
        for k, v in self._old.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        cap._CACHE.clear()
        self._tmp.cleanup()

    def _fake_install(self, *, scripts=_REQ) -> Path:
        """造一个"装好了"的 ppt-master 骨架：脚本齐（内容无所谓）+ 自检能过。"""
        home = Path(os.environ["PPT_MASTER_HOME"])
        sd = home / "skills" / "ppt-master" / "scripts"
        sd.mkdir(parents=True, exist_ok=True)
        for name in scripts:
            (sd / name).write_text("# stub\n", encoding="utf-8")
        return home


class TestStatus(_Sandbox):
    """安装自检：没装 / 装了一半 / 装好了，三种都要给出可读结论且不抛。"""

    def test_not_installed(self):
        st = pp.status()
        self.assertFalse(st["installed"])
        self.assertFalse(st["intake_ready"])
        self.assertEqual(sorted(st["missing_required_scripts"]), sorted(_REQ))

    def test_partial_install(self):
        self._fake_install(scripts=("pptx_intake.py",))
        st = pp.status()
        self.assertTrue(st["installed"])
        self.assertFalse(st["intake_ready"], "脚本不齐就不该说能用")
        self.assertIn("beautify_inventory.py", st["missing_required_scripts"])

    def test_ready_install(self):
        self._fake_install()
        st = pp.status()
        self.assertTrue(st["intake_ready"])
        self.assertEqual(st["missing_required_scripts"], [])


class TestCapability(_Sandbox):
    """能力开关：默认关、关着就不 ready、字段要进快照与 /api/mcp/status。"""

    def test_registered_and_default_off(self):
        ids = [s["id"] for s in cap.snapshot()["skills"]]
        self.assertIn("ppt-polish", ids)
        row = next(s for s in cap.snapshot()["skills"] if s["id"] == "ppt-polish")
        self.assertTrue(row["builtin"], "内置能力，UI 不该显示「文件缺失」")
        self.assertFalse(cap.ppt_polish_enabled(), "**默认关闭**：它会改写老师的稿，必须显式开")
        self.assertFalse(cap.ppt_polish_ready())

    def test_ready_requires_switch_and_install(self):
        self._fake_install()
        self.assertFalse(cap.ppt_polish_ready(), "没开开关就不该 ready")
        cap.set_enabled("skill", "ppt-polish", True)
        self.assertTrue(cap.ppt_polish_ready(), "开关开 + 装好了 = ready")
        cap.set_enabled("skill", "ppt-polish", False)

    def test_switch_does_not_touch_wps(self):
        """★ 解耦：开/关 PPT 优化都不该动 WPS 与预览开关。"""
        before = {k: cap.is_enabled(*k) for k in (("mcp", "wps-office"),
                                                  ("skill", "doc-session-preview"))}
        cap.set_enabled("skill", "ppt-polish", True)
        after = {k: cap.is_enabled(*k) for k in before}
        self.assertEqual(before, after)
        cap.set_enabled("skill", "ppt-polish", False)

    def test_status_field_exposed(self):
        self.assertIn("ppt_polish_ready", cap.snapshot())
        self.assertIn("ppt_polish_ready", ws.api_mcp_status())


class TestGate(_Sandbox):
    """门禁：开关关 → 403；没装 → 424；都过 → 才干活。"""

    def test_off_is_403(self):
        r = ws.api_ppt_polish(ws.PptPolishReq(job_id="x", actions=["font"]))
        self.assertEqual(getattr(r, "status_code", 200), 403)
        self.assertEqual(json.loads(r.body)["gate"], "capability")

    def test_not_installed_is_424(self):
        cap.set_enabled("skill", "ppt-polish", True)
        r = ws.api_ppt_polish(ws.PptPolishReq(job_id="x", actions=["font"]))
        self.assertEqual(getattr(r, "status_code", 200), 424)
        self.assertEqual(json.loads(r.body)["gate"], "install")

    def test_status_endpoint_never_gated(self):
        """面板要靠它决定按钮可不可点，所以没开/没装也得能给结论。"""
        d = ws.api_ppt_status()
        self.assertTrue(d["ok"])
        self.assertFalse(d["enabled"])
        self.assertIn("actions", d)
        self.assertEqual([a["id"] for a in d["actions"]], ["font", "size", "strip-anim"])

    def test_download_requires_result(self):
        r = ws.api_ppt_download(job_id="nope")
        self.assertEqual(getattr(r, "status_code", 200), 404)


class TestJobDir(_Sandbox):
    """任务目录：原稿只读副本、可清理、过期能自清，且**绝不碰 content/plans**。"""

    def setUp(self):
        super().setUp()
        self.src = Path(self._tmp.name) / "老师原稿.pptx"
        _make_pptx(self.src)
        self._before = self.src.read_bytes()

    def test_new_job_copies_and_leaves_source(self):
        job = pp.new_job(self.src)
        jd = Path(job["dir"])
        self.assertTrue((jd / "source.pptx").is_file())
        self.assertEqual(self.src.read_bytes(), self._before, "原稿必须一字未改")

    def test_cleanup_job(self):
        job = pp.new_job(self.src)
        pp.cleanup_job(job["job_id"])
        self.assertFalse(Path(job["dir"]).exists())

    def test_cleanup_stale_only_old(self):
        old, fresh = pp.new_job(self.src), pp.new_job(self.src)
        oldp = Path(old["dir"])
        os.utime(oldp, (time.time() - 48 * 3600,) * 2)
        n = pp.cleanup_stale(24)
        self.assertGreaterEqual(n, 1)
        self.assertFalse(oldp.exists())
        self.assertTrue(Path(fresh["dir"]).exists(), "新的不能被清")

    def test_never_touches_content_plans(self):
        plans = _PROJ / "content" / "plans"
        before = {p.name for p in plans.glob("*")} if plans.is_dir() else set()
        job = pp.new_job(self.src)
        pp.polish(job, ["font", "size", "strip-anim"])
        pp.cleanup_job(job["job_id"])
        after = {p.name for p in plans.glob("*")} if plans.is_dir() else set()
        self.assertEqual(before, after, "PPT 优化不得在 content/plans 落任何东西")


class TestPolish(_Sandbox):
    """保守原地统一：字体 / 字号层级 / 清动画；输出是新文件，原稿不动。"""

    def setUp(self):
        super().setUp()
        self.src = Path(self._tmp.name) / "deck.pptx"
        _make_pptx(self.src, with_anim=True)
        self._before = self.src.read_bytes()
        self.job = pp.new_job(self.src)

    def test_unify_font(self):
        rep = pp.polish(self.job, ["font"], target_ea="微软雅黑", target_latin="Calibri")
        self.assertTrue(rep["ok"], rep)
        self.assertIn("runs", rep["applied"][0])
        self.assertGreater(rep["applied"][0]["runs"], 0)
        from pptx import Presentation
        prs = Presentation(rep["file"])
        names = {r.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs}
        self.assertEqual(names, {"Calibri"}, f"西文字体应统一，实际 {names}")

    def test_unify_size_snapshots_to_anchors(self):
        rep = pp.polish(self.job, ["size"], size_anchors=[44, 32, 24, 12])
        self.assertTrue(rep["ok"], rep)
        from pptx import Presentation
        prs = Presentation(rep["file"])
        sizes = {round(float(r.font.size.pt), 1) for s in prs.slides for sh in s.shapes
                 if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs
                 if r.font.size is not None}
        self.assertTrue(sizes.issubset({44.0, 32.0, 24.0, 12.0}), f"字号应落到锚点，实际 {sizes}")

    def test_strip_anim_removes_timing(self):
        rep = pp.polish(self.job, ["strip-anim"])
        self.assertTrue(rep["ok"], rep)
        self.assertGreaterEqual(rep["applied"][0]["timing_removed"], 1, "造了一个 timing 节点，应删掉")
        import zipfile
        with zipfile.ZipFile(rep["file"]) as z:
            xml = z.read("ppt/slides/slide2.xml").decode("utf-8", "replace")
        self.assertNotIn("<p:timing", xml, "动画节点应被清掉")

    def test_source_untouched_and_result_is_new_file(self):
        rep = pp.polish(self.job, ["font"])
        self.assertNotEqual(Path(rep["file"]), self.src, "产物必须是新文件")
        self.assertEqual(self.src.read_bytes(), self._before, "原稿必须一字未改")

    def test_empty_actions_is_refused(self):
        rep = pp.polish(self.job, [])
        self.assertFalse(rep["ok"])
        self.assertIn("font", rep["available"])


class TestIntakeParse(_Sandbox):
    """体检报告解析：用假脚本喂固定 JSON，验我们提炼的字段（不依赖真安装）。"""

    def setUp(self):
        super().setUp()
        home = self._fake_install()
        sd = home / "skills" / "ppt-master" / "scripts"
        # 假 pptx_intake：写出三份产物（形状照着真脚本的关键字段来）
        (sd / "pptx_intake.py").write_text(
            "import json,sys,pathlib\n"
            "o=pathlib.Path(sys.argv[sys.argv.index('-o')+1]); o.mkdir(parents=True,exist_ok=True)\n"
            "stem=pathlib.Path(sys.argv[-1]).stem\n"
            "(o/'source_profile.json').write_text(json.dumps({'schema':'x','decks':[{'slide_count':3,"
            "'canvas':{'width_px':1280,'height_px':720}}]},ensure_ascii=False),encoding='utf-8')\n"
            "(o/(stem+'.slide_library.json')).write_text('{}',encoding='utf-8')\n"
            "(o/'source.identity.json').write_text(json.dumps({'theme':{'fonts':{'title':{'latin':'Calibri','ea':'微软雅黑'},"
            "'body':{'latin':'Calibri','scripts':{'Hans':'宋体'}}},'sizes':{'title':44.0,'body':32.0,'body_levels':[32,28,24]},"
            "'palette':{'primary':'#4F81BD'}},'observed':{'fonts':{'latin':['Arial'],'ea':['宋体']},"
            "'sizes_pt':[{'value':20.0,'chars':120}],'colors':['#FF0000']}},ensure_ascii=False),encoding='utf-8')\n",
            encoding="utf-8")
        # 假 beautify_inventory：--summary 输出台账、builder 模式写文件
        (sd / "beautify_inventory.py").write_text(
            "import json,sys\n"
            "if '--summary' in sys.argv:\n"
            "    print(json.dumps({'schema':'beautify_inventory.view.v1','slide_count':3,'slides':[\n"
            "      {'slide_index':1,'page_type':'cover_candidate','text_block_count':2,'text_char_count':12,"
            "'table_count':0,'chart_count':0,'diagram_count':0,'image_count':0,'ignored':[],'needs_confirmation':[]},\n"
            "      {'slide_index':2,'page_type':'content_candidate','text_block_count':3,'text_char_count':180,"
            "'table_count':1,'chart_count':1,'diagram_count':0,'image_count':2,'ignored':['hidden shape'],"
            "'needs_confirmation':['merged table']},\n"
            "      {'slide_index':3,'page_type':'ending_candidate','text_block_count':1,'text_char_count':8,"
            "'table_count':0,'chart_count':0,'diagram_count':0,'image_count':0,'ignored':[],'needs_confirmation':[]}]},"
            "ensure_ascii=False))\n"
            "else:\n"
            "    p=sys.argv[sys.argv.index('-o')+1]; open(p,'w',encoding='utf-8').write('{}')\n"
            "    print('written')\n",
            encoding="utf-8")
        self.src = Path(self._tmp.name) / "deck.pptx"
        _make_pptx(self.src)

    def test_report_fields(self):
        rep = pp.intake(pp.new_job(self.src))
        self.assertTrue(rep["ok"], rep)
        self.assertEqual(rep["slides"], 3)
        self.assertEqual(rep["theme"]["title_font"], "微软雅黑", "有 ea 就该取 ea（中文字体）")
        self.assertEqual(rep["theme"]["body_font"], "宋体", "ea 空时退到 scripts.Hans")
        self.assertEqual(rep["theme"]["body_size"], 32.0)
        self.assertEqual(rep["totals"]["tables"], 1)
        self.assertEqual(rep["totals"]["charts"], 1)
        self.assertEqual(rep["totals"]["images"], 2)
        self.assertEqual([p["page"] for p in rep["pages"]], [1, 2, 3])
        self.assertEqual(rep["pages"][1]["confirm"], ["merged table"])
        self.assertIn(2, rep["flags"]["needs_confirmation"])
        self.assertIn(3, rep["flags"]["sparse_pages"])
        self.assertTrue(Path(rep["artifacts"]["inventory"]).is_file())

    def test_intake_without_install_is_explicit(self):
        os.environ["PPT_MASTER_HOME"] = str(Path(self._tmp.name) / "nope")
        rep = pp.intake({"job_id": "x", "dir": self._tmp.name, "source": str(self.src), "name": "deck.pptx"})
        self.assertFalse(rep["ok"])
        self.assertIn("未就绪", rep["error"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
