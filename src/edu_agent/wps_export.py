# -*- coding: utf-8 -*-
"""教案 → WPS 文字（.docx）导出（2026-09-10 WPS 集成）。

链路：
    教案中心 DZ 模型 / Agent B Markdown
      -> 组装 MCP 调用序列（对齐 wps-office-mcp 工具契约）
      -> wps_common_ping -> wps_execute_method(createDocument) -> wps_word_insert_text ×N
      -> wps_execute_method(insertTable) 汇总表 -> wps_common_save_as(.docx)
      -> content/plans/<课题>.docx

闸门与「MCP 服务 / Skill」开关**同源**：capabilities 里 wps-office 关掉就直接拒绝，
不起子进程、不碰 WPS。

工具契约要点（静态核对 mcp_servers/wps-office-mcp）：
- 没有现成的「Markdown/HTML -> 文档」工具，必须自己拆块后逐条 insert_text
- insert_text 的 style 施加在 Selection.Range 上，故 position 固定用 cursor
- createDocument 必须先于所有插入（其余动作都取 ActiveDocument）
- saveAs 是扩展名驱动；Word 分支不会自动建父目录，父目录要自己 mkdir
"""
from __future__ import annotations

import os
import re
from pathlib import Path

from edu_agent import capabilities, mcp_tools

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "content" / "plans"

SERVICE = "wps-office"

# PPT 版式（桥接 addSlide 的 layout 键）
PPT_TITLE_LAYOUT = "title"
PPT_BODY_LAYOUT = "title_content"
# PPT 正文换行：用 CR 才是「新段落（带项目符号）」，LF 只会变成段内换行
PPT_BR = "\r"

# 中文 Word 样式名（缺失样式在桥接里会被 try/catch 吞掉，不会中断导出）
STYLE_TITLE = "标题 1"
STYLE_H2 = "标题 2"
STYLE_H3 = "标题 3"
STYLE_BODY = "正文"
STYLE_BULLET = "列表段落"


# ---------- 组装调用序列 ----------
def _insert(text: str, style: str = STYLE_BODY) -> tuple[str, dict]:
    return ("wps_word_insert_text",
            {"text": text, "position": "cursor", "style": style, "new_paragraph": True})


def _create_document() -> tuple[str, dict]:
    return ("wps_execute_method", {"method": "createDocument", "params": {}})


def _table(grid: list[list[str]]) -> tuple[str, dict] | None:
    if not grid:
        return None
    cols = max(len(r) for r in grid)
    norm = [list(r) + [""] * (cols - len(r)) for r in grid]
    return ("wps_execute_method",
            {"method": "insertTable",
             "params": {"rows": len(norm), "cols": cols, "data": norm}})


def _save_as(path: Path) -> tuple[str, dict]:
    return ("wps_common_save_as", {"filePath": str(path), "format": "docx"})


def _intro(title: str, subtitle: str) -> list[tuple[str, dict]]:
    calls = [_insert(title or "教案", STYLE_TITLE)]
    if subtitle:
        calls.append(_insert(subtitle, STYLE_BODY))
    return calls


# ---------- 入口 1：教案中心 DZ 模型 ----------
def design_calls(payload: dict, out_path: Path) -> list[tuple[str, dict]]:
    """教案中心「教案 × PPT 编排」的数据模型 -> 调用序列。"""
    title = str(payload.get("title") or "").strip() or "未命名课题"
    grade = str(payload.get("grade") or "").strip()
    steps = payload.get("steps") or []
    steps = [s for s in steps if isinstance(s, dict)]

    total = 0
    for s in steps:
        try:
            total += int(s.get("minutes") or 0)
        except (TypeError, ValueError):
            pass

    sub = grade
    if steps:
        sub = (grade + "　·　" if grade else "") + f"共 {len(steps)} 个环节 / {total} 分钟"

    calls = [_create_document(), *_intro(title, sub), _insert("一、教案", STYLE_H2)]

    grid: list[list[str]] = [["环节", "时长", "要点"]]
    for i, s in enumerate(steps, 1):
        name = str(s.get("name") or f"环节 {i}").strip()
        mins = str(s.get("minutes") or "").strip()
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        calls.append(_insert(f"{i}. {name}（{mins or '—'} 分钟）", STYLE_H3))
        if bullets:
            for b in bullets:
                calls.append(_insert(b, STYLE_BULLET))
        else:
            calls.append(_insert("（未填要点）", STYLE_BODY))
        grid.append([f"{i}. {name}", (mins + " 分钟") if mins else "—",
                     "；".join(bullets[:3])[:120] or "—"])

    if steps:
        calls.append(_insert("二、环节一览", STYLE_H2))
        t = _table(grid)
        if t:
            calls.append(t)

    calls.append(_insert("三、PPT 大纲（每页对应一个环节）", STYLE_H2))
    for i, s in enumerate(steps, 1):
        name = str(s.get("name") or f"环节 {i}").strip()
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        calls.append(_insert(f"第 {i} 页　{name}", STYLE_H3))
        for b in bullets[:5]:
            calls.append(_insert(b, STYLE_BULLET))
    if not steps:
        calls.append(_insert("（尚无环节，先在教案中心生成骨架）", STYLE_BODY))

    calls.append(_insert("导出自 edu-agent 教案中心 · WPS 文字", STYLE_BODY))
    calls.append(_save_as(out_path))
    return calls


# ---------- 入口 2：Agent B 的 Markdown 教案 ----------
_H_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_LI_RE = re.compile(r"^\s*(?:[-*+]|\d+[.、)])\s+(.*)$")


def _flush_table(rows: list[list[str]], calls: list[tuple[str, dict]]) -> None:
    if len(rows) >= 2:
        body = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r if c.strip())]
        t = _table(body)
        if t:
            calls.append(t)
    elif rows:
        for r in rows:
            calls.append(_insert("　".join(r), STYLE_BODY))


def markdown_calls(markdown_text: str, title: str, out_path: Path) -> list[tuple[str, dict]]:
    """Agent B 生成的 Markdown 教案 -> 调用序列（标题/正文/列表/表格）。"""
    calls = [_create_document(), *_intro(title or "教案", "依据教材生成 · 导出自 edu-agent")]
    tbl: list[list[str]] = []
    for raw in (markdown_text or "").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("|") and line.strip().endswith("|"):
            tbl.append([c.strip() for c in line.strip().strip("|").split("|")])
            continue
        if tbl:
            _flush_table(tbl, calls)
            tbl = []
        if not line.strip():
            continue
        m = _H_RE.match(line)
        if m:
            level = len(m.group(1))
            style = STYLE_TITLE if level == 1 else (STYLE_H2 if level == 2 else STYLE_H3)
            calls.append(_insert(m.group(2).strip(), style))
            continue
        m = _LI_RE.match(line)
        if m:
            calls.append(_insert(m.group(1).strip(), STYLE_BULLET))
            continue
        calls.append(_insert(re.sub(r"\*\*(.+?)\*\*", r"\1", line.strip()), STYLE_BODY))
    if tbl:
        _flush_table(tbl, calls)
    calls.append(_save_as(out_path))
    return calls


# ---------- 落盘与执行 ----------
def _safe_name(title: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", title or "教案")[:50] or "教案"


def resolve_out_dir(out_dir: str | None) -> Path:
    """解析导出目录：空 -> content/plans；相对路径按项目根；不存在则创建；失败回落默认。"""
    if not out_dir:
        return OUT_DIR
    try:
        raw = os.path.expandvars(str(out_dir).strip().strip('"').strip("'"))
        p = Path(raw).expanduser()
        if not p.is_absolute():
            p = ROOT / p
        p.mkdir(parents=True, exist_ok=True)            # saveAs 不会建父目录
        return p if p.is_dir() else OUT_DIR
    except Exception:
        return OUT_DIR


def _payload_title(payload: dict, markdown: str = "") -> str:
    title = str(payload.get("title") or "").strip()
    if not title and markdown:
        m = _H_RE.match(next((l for l in markdown.splitlines() if l.strip()), ""))
        title = m.group(2).strip() if m else ""
    return title or "未命名课题"


# ---------- PPT（WPS 演示 .pptx） ----------
def pptx_native(title: str, grade: str, steps: list[dict], out_path: Path,
                markdown: str = "") -> dict:
    """用 python-pptx 本地生成 .pptx（不依赖 WPS COM）。

    2026-09-10 决策：本机 WPS 演示的 COM 通道是「无头代理进程」（wpp.exe MainWindowHandle=0），
    Presentations.Add 出来的文稿 Windows/Slides 恒为 0，Slides.Add 返回 null，SaveAs 报
    RPC_E_CALL_REJECTED —— 即 WPS 演示这条 COM 路在本机不可用。故 PPT 改为本地生成真 OOXML，
    生成后再尽力用 WPS 打开给用户看（见 open_with_wps）。
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)          # 16:9
    prs.slide_height = Inches(7.5)

    def _body_ph(slide):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                return ph
        return None

    # 封面
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    sub = (grade + "　" if grade else "")
    total = 0
    for s in steps:
        try:
            total += int(s.get("minutes") or 0)
        except (TypeError, ValueError):
            pass
    if steps:
        sub += f"共 {len(steps)} 个环节 / {total} 分钟"
    body = _body_ph(slide)
    if body is not None and sub:
        body.text = sub
        for p in body.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)

    # 每环节一页
    for i, s in enumerate(steps, 1):
        name = str(s.get("name") or f"环节 {i}").strip()
        mins = str(s.get("minutes") or "").strip()
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = f"{i}. {name}" + (f"（{mins} 分钟）" if mins else "")
        ph = _body_ph(slide)
        if ph is not None:
            tf = ph.text_frame
            tf.clear()
            items = bullets[:6] or ["（未填要点）"]
            for j, b in enumerate(items):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                para.text = b
                for r in para.runs:
                    r.font.size = Pt(20)

    if not steps:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = "（尚无环节）"
        ph = _body_ph(slide)
        if ph is not None:
            ph.text_frame.text = "先在教案中心生成骨架或添加环节"

    # Markdown 源的要点页（Agent B 教案）
    if markdown.strip() and not steps:
        cur = None
        for raw in markdown.splitlines():
            m = _H_RE.match(raw.rstrip())
            if m and len(m.group(1)) == 2:
                cur = m.group(2).strip()
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                if slide.shapes.title is not None:
                    slide.shapes.title.text = cur
                ph = _body_ph(slide)
                if ph is not None:
                    ph.text_frame.clear()
            elif cur:
                li = _LI_RE.match(raw)
                slide = prs.slides[-1]
                ph = _body_ph(slide)
                if li and ph is not None:
                    tf = ph.text_frame
                    para = tf.paragraphs[0] if not tf.text.strip() else tf.add_paragraph()
                    para.text = li.group(1).strip()
                    for r in para.runs:
                        r.font.size = Pt(20)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {"slides": len(prs.slides.__iter__.__self__._sldIdLst)}


def open_with_wps(path: Path, timeout: float = 20.0) -> str:
    """生成后尽力用 WPS 打开（先试 WPS MCP，失败回落系统默认程序）。返回打开方式。"""
    if capabilities.is_enabled("mcp", "wps-office"):
        tool = ("wps_ppt_open_presentation" if path.suffix.lower() == ".pptx"
                else "wps_word_open_document")
        r = mcp_tools.call_tool(SERVICE, tool, {"filePath": str(path)}, timeout=timeout)
        if r:
            return "wps-mcp"
    try:
        os.startfile(str(path))          # noqa: S606 - 本机单用户，用默认程序打开导出物
        return "os-startfile"
    except Exception:
        return "none"


# ---------- 本地 .docx 生成（WPS COM 幽灵态时的兜底） ----------
def _set_east_asian(style, ascii_font: str, ea_font: str, size_pt: float | None = None) -> None:
    """给样式设中英文字体（python-docx 不直接暴露 eastAsia，需走 rPr）。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = ascii_font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), ea_font)
    if size_pt:
        style.font.size = Pt(size_pt)


def docx_native(title: str, grade: str, steps: list[dict], out_path: Path,
                markdown: str = "") -> dict:
    """用 python-docx 本地生成 .docx（结构与 WPS COM 版教案一致）。

    为什么需要兜底：本机 WPS 11.1.0.10009 的 Writer COM 会话会进入「幽灵态」——
    Documents.Add() 建出空壳文档（path=null、0 字符），insertText 报成功却不落字，
    SaveAs 报「另存为成功」却不写盘（强杀 WPS / 删掉它打开的目录后容易触发）。
    因此导出链路是「WPS COM 优先，未落盘则本地兜底」，保证交付物一定产出。
    """
    from docx import Document

    doc = Document()
    _set_east_asian(doc.styles["Normal"], "Times New Roman", "宋体", 12)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12.5),
                       ("List Bullet", 12)):
        try:
            _set_east_asian(doc.styles[name], "Times New Roman", "黑体", size)
        except KeyError:
            pass

    doc.add_heading(title or "未命名课题", level=1)

    total = 0
    for s in steps:
        try:
            total += int(s.get("minutes") or 0)
        except (TypeError, ValueError):
            pass
    meta = (grade + "　" if grade else "")
    if steps:
        meta += f"共 {len(steps)} 个环节 / {total} 分钟"
    if meta.strip():
        doc.add_paragraph(meta)

    doc.add_heading("一、教案", level=2)
    for i, s in enumerate(steps, 1):
        name = str(s.get("name") or f"环节 {i}").strip()
        mins = str(s.get("minutes") or "").strip()
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        doc.add_heading(f"{i}. {name}" + (f"（{mins} 分钟）" if mins else ""), level=3)
        if bullets:
            for b in bullets:
                doc.add_paragraph(b, style="List Bullet")
        else:
            doc.add_paragraph("（未填要点）")

    if not steps and (markdown or "").strip():
        # Markdown 源（Agent B 教案）：H2 当小节标题
        for raw in markdown.splitlines():
            m = _H_RE.match(raw.rstrip())
            if m and len(m.group(1)) == 2:
                doc.add_heading(m.group(2).strip(), level=3)
                continue
            li = _LI_RE.match(raw)
            if li:
                doc.add_paragraph(li.group(1).strip(), style="List Bullet")
            elif raw.strip():
                doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", raw.strip()))

    if steps:
        doc.add_heading("二、环节一览", level=2)
        t = doc.add_table(rows=1, cols=3)
        try:
            t.style = "Table Grid"
        except KeyError:
            pass
        for cell, txt in zip(t.rows[0].cells, ("环节", "时长", "要点")):
            cell.text = txt
        for i, s in enumerate(steps, 1):
            cells = t.add_row().cells
            cells[0].text = f"{i}. {str(s.get('name') or '').strip()}"
            cells[1].text = (str(s.get("minutes") or "").strip() + " 分钟").strip()
            bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
            cells[2].text = "；".join(bullets[:3])[:120] or "—"

        doc.add_heading("三、PPT 大纲（每页对应一个环节）", level=2)
        for i, s in enumerate(steps, 1):
            doc.add_heading(f"第 {i} 页　{str(s.get('name') or '').strip()}", level=3)
            bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
            for b in bullets[:5]:
                doc.add_paragraph(b, style="List Bullet")

    doc.add_paragraph("导出自 edu-agent 教案中心")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"paragraphs": len(doc.paragraphs)}


def export_docx(payload: dict, markdown: str = "", out_dir: str | None = None) -> dict:
    """Word 导出（保留旧签名，内部走 export）。"""
    return export(payload, markdown=markdown, kind="docx", out_dir=out_dir)


# ---------- 富模板（用户导入并编辑过的教案模板）直出 ----------
def _set_run_east_asian(run, font: str) -> None:
    """给单个 run 设中文字体（eastAsia 属性，python-docx 不直接暴露）。"""
    from docx.oxml.ns import qn

    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _apply_runs(par, runs: list[dict], fallback_text: str = "") -> None:
    from docx.shared import Pt, RGBColor

    items = runs or [{"text": fallback_text}]
    if not any((r.get("text") or "").strip() for r in items):
        items = [{"text": fallback_text}]
    for r in items:
        run = par.add_run(str(r.get("text") or ""))
        if r.get("b"):
            run.bold = True
        if r.get("i"):
            run.italic = True
        if r.get("u"):
            run.underline = True
        if r.get("font"):
            _set_run_east_asian(run, str(r["font"]))
        if r.get("size"):
            try:
                run.font.size = Pt(float(r["size"]))
            except Exception:
                pass
        col = str(r.get("color") or "").lstrip("#")
        if len(col) == 6:
            try:
                run.font.color.rgb = RGBColor.from_string(col.upper())
            except Exception:
                pass


def _apply_align(par, align: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    m = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    if align in m:
        par.alignment = m[align]


def docx_from_rich(rich: dict, out_path: Path) -> dict:
    """富模板 -> .docx（逐 run 还原字体/字号/加粗/斜体/下划线/颜色 + 段落对齐 + 表格）。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _set_east_asian(doc.styles["Normal"], "Times New Roman", "宋体", 12)
    blocks = [b for b in (rich.get("blocks") or []) if isinstance(b, dict)]
    skipped_images = 0
    tables = 0
    paras = 0

    for bi, b in enumerate(blocks):
        title = str(b.get("title") or "").strip()
        if title:
            par = doc.add_paragraph()
            _apply_align(par, b.get("title_align") or "")
            runs = b.get("title_runs") or []
            if runs:
                _apply_runs(par, runs, title)
            else:                       # 新建板块：给个体面的默认标题样式
                run = par.add_run(title)
                run.bold = True
                run.font.size = Pt(15)
                _set_run_east_asian(run, "黑体")
            if bi > 0:
                par.paragraph_format.space_before = Pt(10)
        for el in (b.get("elements") or []):
            et = el.get("type")
            if et == "para":
                par = doc.add_paragraph()
                _apply_align(par, el.get("align") or "")
                ls = el.get("line_spacing")
                if ls:
                    try:
                        if float(ls) > 3:            # 固定值（磅）
                            par.paragraph_format.line_spacing = Pt(float(ls))
                        else:
                            par.paragraph_format.line_spacing = float(ls)
                    except Exception:
                        pass
                _apply_runs(par, el.get("runs") or [], str(el.get("text") or ""))
                paras += 1
            elif et == "table":
                cells = el.get("cells") or []
                if not cells:
                    continue
                cols = max(len(r) for r in cells)
                t = doc.add_table(rows=len(cells), cols=cols)
                try:
                    t.style = "Table Grid"
                except KeyError:
                    pass
                for ri, row in enumerate(cells):
                    for ci in range(cols):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        data = row[ci] if ci < len(row) else {}
                        if isinstance(data, dict):
                            txt = str(data.get("text") or "")
                            if data.get("b") or (el.get("header", True) and ri == 0):
                                run = p.add_run(txt)
                                run.bold = True
                            else:
                                p.add_run(txt)
                            _apply_align(p, data.get("align") or ("center" if ri == 0 else ""))
                        else:
                            p.add_run(str(data))
                tables += 1
            elif et == "image":
                skipped_images += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"paragraphs": paras, "tables": tables, "images_skipped": skipped_images}


def pptx_from_rich(rich: dict, out_path: Path) -> dict:
    """富模板 -> .pptx：一个板块一页（标题=板块名，正文=该板块前几段文字）。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def body_ph(slide):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                return ph
        return None

    def add(title: str, lines: list[str]):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = title[:80]
        ph = body_ph(slide)
        if ph is not None:
            tf = ph.text_frame
            tf.clear()
            for i, ln in enumerate(lines[:8] or ["（本板块暂无内容）"]):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = ln[:120]
                for r in para.runs:
                    r.font.size = Pt(18)

    # 封面
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title is not None:
        slide.shapes.title.text = str(rich.get("name") or "教案")
    sub = body_ph(slide)
    if sub is not None:
        sub.text = "依据导入模板生成 · edu-agent 教案中心"

    n = 1
    for b in (rich.get("blocks") or []):
        if b.get("kind") == "preamble":
            continue
        lines = [str(e.get("text") or "") for e in (b.get("elements") or [])
                 if e.get("type") == "para" and (e.get("text") or "").strip()]
        for e in (b.get("elements") or []):
            if e.get("type") == "table":
                rows = e.get("cells") or []
                if rows:
                    lines.append("表格：" + " / ".join(str(c.get("text") or "") for c in rows[0])[:80])
                break
        add(str(b.get("title") or f"板块{n}"), lines)
        n += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {"slides": n}


def _export_from_template(template: dict, title: str, out: Path, kind: str,
                          open_file: bool = True) -> dict:
    """富模板直出（docx / pptx），返回结构与 export() 一致。"""
    import time

    from edu_agent import template_rich

    tpl = template_rich.normalize(template)
    out_path = out / (_safe_name(title or tpl.get("name") or "教案") + "." + kind)
    t0 = time.time()
    try:
        info = (pptx_from_rich(tpl, out_path) if kind == "pptx" else docx_from_rich(tpl, out_path))
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"模板导出失败：{type(e).__name__}: {e}", "gate": "local"}

    if not out_path.exists():
        return {"ok": False, "error": "模板导出未生成文件", "gate": "local"}

    try:
        fmt_ok: bool | None = out_path.read_bytes()[:2] == b"PK"
    except OSError:
        fmt_ok = None
    opened = open_with_wps(out_path) if open_file else "none"
    eng = "python-pptx（模板直出）" if kind == "pptx" else "python-docx（模板直出）"
    skipped = info.get("images_skipped", 0)
    warn = f"（跳过 {skipped} 张图片：导入时不保留图片）" if skipped else ""
    return {"ok": True, "path": str(out_path), "filename": out_path.name, "kind": kind,
            "outDir": str(out), "engine": eng, "opened": opened, "format_ok": fmt_ok,
            "template_id": tpl.get("template_id"), "stats": template_rich.stats(tpl),
            "elapsed": round(time.time() - t0, 1),
            "url": "/files/" + out_path.name,
            "download_url": "/api/wps/download?path=" + str(out_path).replace("\\", "/"),
            "detail": f"已按模板「{tpl.get('name')}」导出 {out_path.name}（{eng}）{warn}"}


def export_pptx(payload: dict, markdown: str = "", out_dir: str | None = None) -> dict:
    """PPT 导出。"""
    return export(payload, markdown=markdown, kind="pptx", out_dir=out_dir)


def export(payload: dict, markdown: str = "", kind: str = "docx",
           out_dir: str | None = None, template: dict | None = None) -> dict:
    """统一导出入口。

    kind: docx（WPS 文字，走 WPS COM）| pptx（python-pptx 本地生成 + 尽力用 WPS 打开）
    out_dir: 自选导出目录（空 = content/plans）
    返回 {ok, path, filename, kind, outDir, calls/failed 或 slides/engine, format_ok, download_url, detail}
    """
    ok, why = capabilities.wps_available()
    if not ok:
        return {"ok": False, "error": why, "gate": "capability"}

    kind = "pptx" if str(kind).lower() in ("pptx", "ppt") else "docx"
    out = resolve_out_dir(out_dir)

    # ---------- 富模板直出（用户导入并编辑过的教案模板） ----------
    if template:
        tpl_title = str(payload.get("title") or "").strip() or str(template.get("name") or "")
        return _export_from_template(template, tpl_title, out, kind,
                                     open_file=payload.get("open", True) is not False)

    title = _payload_title(payload, markdown)

    if kind == "pptx":
        return _export_pptx(payload, markdown, title, out)

    # ---------- Word：WPS COM 优先（真·集成路径），未落盘则 python-docx 本地兜底 ----------
    out_path = out / (_safe_name(title) + ".docx")
    calls = (markdown_calls(markdown, title, out_path) if markdown.strip()
             else design_calls(payload, out_path))

    results = mcp_tools.call_sequence(SERVICE, calls, timeout=420.0)
    if not results:
        return {"ok": False, "error": "WPS MCP 未响应（服务被关闭或桥接失败）", "gate": "mcp"}

    failed = [r for r in results if not r.get("ok")]
    engine = "wps-mcp"
    if not out_path.exists():
        # WPS Writer 会话「幽灵态」：调用报成功却不落盘（见 docx_native 说明）→ 本地兜底
        steps = [s for s in (payload.get("steps") or []) if isinstance(s, dict)]
        grade = str(payload.get("grade") or "").strip()
        try:
            docx_native(title, grade, steps, out_path, markdown=markdown)
            engine = "python-docx"
        except Exception as e:  # noqa: BLE001
            tail = (failed[-1].get("text") or failed[-1].get("error")) if failed else ""
            return {"ok": False,
                    "error": (f"导出未生成文件，本地兜底也失败：{type(e).__name__}: {e}"
                              + (f"（WPS 返回：{str(tail)[:120]}）" if tail else "")),
                    "gate": "wps", "calls": len(results), "failed": len(failed)}

    # 格式自检：docx/pptx 都是 OOXML(zip, 头 504b)。WPS 曾把 Word 的 16 落成二进制 .doc，
    # 这里留哨兵：落盘不是 zip 就在返回值里告警（见 docs/08-WPS集成说明.md）。
    fmt_ok: bool | None = None
    try:
        fmt_ok = out_path.read_bytes()[:2] == b"PK"
    except OSError:
        fmt_ok = None  # 文件仍被 WPS 占用，跳过校验
    warn = "" if fmt_ok is not False else "（警告：落盘不是 OOXML zip，请检查桥接 saveAs 格式映射）"

    opened = open_with_wps(out_path) if payload.get("open", True) is not False else "none"
    opened_txt = {"wps-mcp": "已在 WPS 文字中打开", "os-startfile": "已用默认程序打开",
                  "none": "未自动打开"}.get(opened, opened)
    eng_txt = ("WPS COM（MCP %d 步）" % len(results)) if engine == "wps-mcp" else "python-docx 本地兜底"

    return {"ok": True, "path": str(out_path), "filename": out_path.name, "kind": "docx",
            "outDir": str(out), "calls": len(results), "failed": len(failed),
            "engine": engine, "opened": opened, "format_ok": fmt_ok,
            "url": "/files/" + out_path.name,
            "download_url": "/api/wps/download?path=" + str(out_path).replace("\\", "/"),
            "detail": f"已写入 {out_path.name}（{eng_txt} · {opened_txt}）{warn}"}


def _export_pptx(payload: dict, markdown: str, title: str, out: Path) -> dict:
    """PPT：python-pptx 本地生成（不走 WPS COM，见 pptx_native 的说明），再尽力用 WPS 打开。"""
    import time

    out_path = out / (_safe_name(title) + ".pptx")
    steps = [s for s in (payload.get("steps") or []) if isinstance(s, dict)]
    grade = str(payload.get("grade") or "").strip()
    t0 = time.time()
    try:
        info = pptx_native(title, grade, steps, out_path, markdown=markdown)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"PPT 生成失败：{type(e).__name__}: {e}", "gate": "local"}

    if not out_path.exists():
        return {"ok": False, "error": "PPT 未生成文件", "gate": "local"}

    try:
        fmt_ok: bool | None = out_path.read_bytes()[:2] == b"PK"
    except OSError:
        fmt_ok = None

    opened = open_with_wps(out_path) if payload.get("open", True) is not False else "none"
    opened_txt = {"wps-mcp": "已在 WPS 演示中打开", "os-startfile": "已用默认程序打开",
                  "none": "未自动打开（可点「下载到本机」）"}.get(opened, opened)
    pages = info.get("slides", len(steps) + 1)

    return {"ok": True, "path": str(out_path), "filename": out_path.name, "kind": "pptx",
            "outDir": str(out), "slides": pages, "engine": "python-pptx", "opened": opened,
            "format_ok": fmt_ok, "elapsed": round(time.time() - t0, 1),
            "url": "/files/" + out_path.name,
            "download_url": "/api/wps/download?path=" + str(out_path).replace("\\", "/"),
            "detail": f"已生成 {out_path.name}（{pages} 页 · python-pptx 本地生成 · {opened_txt}）"}


if __name__ == "__main__":  # 手工自检：python -m edu_agent.wps_export [docx|pptx] [输出目录]
    import json
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    kind = sys.argv[1] if len(sys.argv) > 1 else "docx"
    out = sys.argv[2] if len(sys.argv) > 2 else ""
    demo = {"title": "3.3 幂函数", "grade": "重点班 45 分钟",
            "steps": [{"name": "情境导入", "minutes": 5, "bullets": ["回顾指数函数", "提出幂函数问题"]},
                      {"name": "概念建构", "minutes": 15, "bullets": ["五个常见幂函数", "图象与性质"]},
                      {"name": "练习巩固", "minutes": 20, "bullets": ["比较大小", "单调性应用"]}]}
    print(json.dumps(export(demo, kind=kind, out_dir=out), ensure_ascii=False, indent=2))
