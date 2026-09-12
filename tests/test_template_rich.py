# -*- coding: utf-8 -*-
"""富教案模板测试（template_rich：.docx <-> 可编辑结构，含格式与表格）。

全离线：自己用 python-docx 造一个带格式与表格的 .docx，跑
解析 -> 规范化 -> 派生(给 Agent B) -> 导出 -> 再解析 的闭环，不碰网络与 WPS。
"""
import sys
import tempfile
import unittest
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import template_rich as tr      # noqa: E402
from edu_agent import template_spec as ts      # noqa: E402
from edu_agent import wps_export as wx         # noqa: E402


def _make_docx(path: Path) -> None:
    """造一个含标题/正文/加粗斜体/表格的教案模板。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_paragraph("一、教学目标")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("知识与技能：")
    r1.bold = True
    r2 = p.add_run("理解函数的概念")
    r2.italic = True
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
    doc.add_paragraph("教学重点：理解函数的概念，掌握函数定义域、对应关系、值域三要素，学会用集合语言刻画函数。")
    doc.add_paragraph("二、教学过程")
    t = doc.add_table(rows=2, cols=3)
    for i, v in enumerate(("环节", "教师活动", "学生活动")):
        c = t.cell(0, i)
        c.text = ""
        run = c.paragraphs[0].add_run(v)
        run.bold = True
    for i, v in enumerate(("导入", "提问", "回答")):
        t.cell(1, i).text = v
    doc.save(str(path))


class TestParse(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.docx = self.dir / "模板.docx"
        _make_docx(self.docx)

    def test_headings_become_blocks(self):
        rich = tr.docx_to_rich(self.docx)
        titles = [b["title"] for b in rich["blocks"]]
        self.assertIn("一、教学目标", titles)
        self.assertIn("二、教学过程", titles)

    def test_long_sentence_not_a_block(self):
        """整句正文（教学重点：…）不能被当成板块标题。"""
        rich = tr.docx_to_rich(self.docx)
        for b in rich["blocks"]:
            self.assertFalse(b["title"].startswith("教学重点："),
                             "长句被误判成板块标题：" + b["title"][:20])

    def test_runs_and_format_preserved(self):
        rich = tr.docx_to_rich(self.docx)
        blk = rich["blocks"][0]
        para = [e for e in blk["elements"] if e["type"] == "para"][0]
        self.assertEqual(para["align"], "center")
        runs = {(r["text"]): r for r in para["runs"]}
        self.assertTrue(runs["知识与技能："]["b"])
        self.assertTrue(runs["理解函数的概念"]["i"])
        self.assertEqual(runs["理解函数的概念"]["size"], 14.0)
        self.assertEqual(runs["理解函数的概念"]["color"].upper(), "#DC2626")

    def test_table_preserved(self):
        rich = tr.docx_to_rich(self.docx)
        tables = [e for b in rich["blocks"] for e in b["elements"] if e["type"] == "table"]
        self.assertEqual(len(tables), 1)
        cells = tables[0]["cells"]
        self.assertEqual(len(cells), 2)
        self.assertEqual(cells[0][0]["text"], "环节")
        self.assertTrue(cells[0][0]["b"], "表头应标为加粗")
        self.assertEqual(cells[1][2]["text"], "回答")

    def test_stats(self):
        st = tr.stats(tr.docx_to_rich(self.docx))
        self.assertGreaterEqual(st["blocks"], 2)
        self.assertEqual(st["tables"], 1)
        self.assertGreater(st["chars"], 20)


class TestNormalizeAndDerive(unittest.TestCase):
    def test_normalize_fills_ids_and_pads_tables(self):
        raw = {"name": "乱数据 模板", "blocks": [
            {"title": "板块A", "elements": [
                {"type": "para", "text": "x"},
                {"type": "table", "cells": [[{"text": "a"}, {"text": "b"}], [{"text": "c"}]]},
            ]},
            "不是字典",
        ]}
        t = tr.normalize(raw)
        self.assertEqual(t["spec_version"], 2)
        self.assertTrue(t["template_id"])
        self.assertTrue(all(b.get("id") for b in t["blocks"]))
        rows = t["blocks"][0]["elements"][1]["cells"]
        self.assertEqual(len(rows[0]), len(rows[1]), "表格行应补齐等宽")
        self.assertEqual(len(t["blocks"]), 1, "非字典板块应被丢弃")

    def test_normalize_clamps_align_and_level(self):
        t = tr.normalize({"name": "x", "blocks": [
            {"title": "T", "level": 99, "elements": [
                {"type": "para", "text": "a", "align": "middle", "level": -3}]}]})
        b = t["blocks"][0]
        self.assertLessEqual(b["level"], tr.MAX_LEVEL)
        self.assertEqual(b["elements"][0]["align"], "", "非法对齐应被清空")
        self.assertEqual(b["elements"][0]["level"], 0)

    def test_derive_simple_for_agent_b(self):
        rich = {"blocks": [
            {"kind": "preamble", "title": "（模板开头）", "elements": [{"type": "para", "text": "开头说明"}]},
            {"kind": "section", "title": "一、教学目标", "elements": [{"type": "para", "text": "目标正文"}]},
            {"kind": "section", "title": "二、教学过程", "elements": [
                {"type": "table", "cells": [[{"text": "环节"}, {"text": "活动"}]]}]},
        ]}
        blocks = tr.derive_simple(rich)
        self.assertEqual([b["type"] for b in blocks], ["preamble", "heading", "table"])
        self.assertEqual(blocks[1]["title"], "一、教学目标")
        self.assertIn("表头", blocks[2]["note"])
        # 能被 template_spec.Block 直接吃下
        for b in blocks:
            ts.Block(**b)

    def test_blank_template(self):
        r = tr.blank_rich("空白")
        self.assertGreaterEqual(len(r["blocks"]), 3)
        self.assertEqual(tr.stats(r)["tables"], 0)


class TestExportRoundTrip(unittest.TestCase):
    """导出后再解析，验证「板块/段落/格式/表格」都还在。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.docx = self.dir / "模板.docx"
        _make_docx(self.docx)
        self.rich = tr.docx_to_rich(self.docx)

    def test_docx_export_roundtrip(self):
        out = self.dir / "out.docx"
        info = wx.docx_from_rich(self.rich, out)
        self.assertTrue(out.exists())
        self.assertEqual(out.read_bytes()[:2], b"PK")
        self.assertGreaterEqual(info["tables"], 1)
        again = tr.docx_to_rich(out)
        t1 = [b["title"] for b in self.rich["blocks"]]
        t2 = [b["title"] for b in again["blocks"]]
        self.assertEqual(t1, t2, "板块标题应在导出/再解析后保持一致")
        runs0 = [e for e in self.rich["blocks"][0]["elements"] if e["type"] == "para"][0]["runs"]
        runs1 = [e for e in again["blocks"][0]["elements"] if e["type"] == "para"][0]["runs"]
        self.assertTrue(any(r["b"] for r in runs1), "加粗应在导出后保留")
        self.assertTrue(any(r["i"] for r in runs1), "斜体应在导出后保留")
        self.assertTrue(any((r["color"] or "").upper() == "#DC2626" for r in runs1), "颜色应保留")

    def test_pptx_export(self):
        out = self.dir / "out.pptx"
        info = wx.pptx_from_rich(self.rich, out)
        self.assertTrue(out.exists())
        self.assertEqual(out.read_bytes()[:2], b"PK")
        self.assertGreaterEqual(info["slides"], 2)      # 封面 + 板块

    def test_export_from_template_entry(self):
        """export(template=...) 直出，且能识别图片跳过。"""
        rich = dict(self.rich)
        rich["blocks"] = list(rich["blocks"]) + [
            {"id": "bz", "kind": "section", "title": "三、插图", "elements": [
                {"id": "ez", "type": "image", "alt": "图片"}]}]
        r = wx._export_from_template(rich, "往返测试", self.dir, "docx", open_file=False)
        self.assertTrue(r["ok"], r.get("error"))
        self.assertEqual(Path(r["path"]).read_bytes()[:2], b"PK")
        self.assertIn("跳过 1 张图片", r["detail"])


class TestTemplateSpecBridge(unittest.TestCase):
    """富模板落盘后，Agent B 的 template_spec 能读到并派生骨架。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.docx = self.dir / "桥接模板.docx"
        _make_docx(self.docx)
        self._old_dir = tr.TEMPLATES_DIR
        tr.TEMPLATES_DIR = Path(self._tmp.name)
        ts.TEMPLATES_DIR = Path(self._tmp.name)
        tr.ORIG_DIR = Path(self._tmp.name) / "orig"

    def tearDown(self):
        tr.TEMPLATES_DIR = self._old_dir
        ts.TEMPLATES_DIR = self._old_dir
        tr.ORIG_DIR = self._old_dir / "orig"

    def test_save_then_agent_b_reads_derived_blocks(self):
        rich = tr.docx_to_rich(self.docx, template_id="桥接", name="桥接")
        tr.save_rich(rich, docx_source=self.docx)
        self.assertTrue(tr.is_rich("桥接"))
        lst = {t["template_id"]: t for t in ts.list_templates()}
        self.assertTrue(lst["桥接"]["rich"], "列表应标记为富模板")
        spec = ts.get_template("桥接")           # Agent B 的读法
        self.assertEqual(len(spec.blocks), len(rich["blocks"]))
        self.assertIn("一、教学目标", ts.template_sketch(spec))


if __name__ == "__main__":
    unittest.main()
